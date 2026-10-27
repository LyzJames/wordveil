"""Untrusted document parsing inside the no-network worker.

Inputs are bytes supplied by the trusted core, never paths or URLs. Nothing in a
document is executed: no macros, scripts, external links, OLE objects or remote
images. Results are bounded and paginated so a single IPC line stays small.

Coordinates are reported in a uniform normalized frame: [x0, y0, x1, y1] with a
top-left origin, relative to the displayed page (after the page's own /Rotate).
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime, time

MAX_DOCUMENT_BYTES = 200 * 1024 * 1024
MAX_TEXT_BYTES = 50 * 1024 * 1024
MAX_PAGES = 2000
MAX_ZIP_UNCOMPRESSED = 512 * 1024 * 1024
MAX_ZIP_RATIO = 200
MAX_ZIP_MEMBERS = 20000
MAX_SHEET_CELLS = 2_000_000
MAX_ROWS_PER_CALL = 2000
OCR_PIXEL_BUDGET = 4_000_000          # below the model adapter's 4,194,304 limit
MIN_NATIVE_CHARS = 12                 # sparse text still requires OCR


class ParseError(ValueError):
    """A document problem the user can act on. `code` is a stable identifier."""

    def __init__(self, code: str, detail: str = ''):
        super().__init__(code)
        self.code = code
        self.detail = detail


def sniff(data: bytes) -> str:
    """Identify the container from magic bytes. Extensions are never trusted."""
    head = data[:16]
    if head.startswith(b'%PDF-'):
        return 'pdf'
    if head.startswith(b'PK\x03\x04'):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
        except zipfile.BadZipFile:
            return 'unknown'
        if 'word/document.xml' in names:
            return 'docx'
        if 'xl/workbook.xml' in names:
            return 'xlsx'
        return 'zip'
    if head.startswith(b'\x89PNG\r\n\x1a\n'):
        return 'image'
    if head[:3] == b'\xff\xd8\xff':
        return 'image'
    if head[:6] in (b'GIF87a', b'GIF89a') or head[:2] == b'BM':
        return 'image'
    if head[:4] in (b'II*\x00', b'MM\x00*'):
        return 'image'
    if head[:4] == b'RIFF' and data[8:12] == b'WEBP':
        return 'image'
    if head.startswith(b'\xd0\xcf\x11\xe0'):
        return 'ole'            # legacy .doc/.xls are not supported in this version
    if decode_text(data[:65536], probe=True) is not None:
        return 'text'
    return 'unknown'


def check_zip(data: bytes) -> None:
    """Reject zip bombs and encrypted members before any XML parser runs."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_MEMBERS:
                raise ParseError('archive_too_many_members')
            total = 0
            for member in members:
                if member.flag_bits & 0x1:
                    raise ParseError('encrypted_document')
                total += member.file_size
                if member.compress_size and member.file_size / max(member.compress_size, 1) > MAX_ZIP_RATIO and member.file_size > 10 * 1024 * 1024:
                    raise ParseError('archive_compression_ratio')
            if total > MAX_ZIP_UNCOMPRESSED:
                raise ParseError('archive_too_large')
            names = {member.filename.lower() for member in members}
            if any(name.endswith('vbaproject.bin') for name in names):
                raise ParseError('macro_document_not_supported')
    except zipfile.BadZipFile as error:
        raise ParseError('corrupt_archive') from error


def decode_text(data: bytes, probe: bool = False):
    """Decode text with an explicit candidate order common for Chinese exports."""
    if data.startswith(b'\xef\xbb\xbf'):
        candidates = ['utf-8-sig']
    elif data.startswith((b'\xff\xfe', b'\xfe\xff')):
        candidates = ['utf-16']
    else:
        candidates = ['utf-8', 'gb18030']
    for encoding in candidates:
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            if probe and encoding == 'utf-8':
                # A probe may cut a multi-byte sequence at the boundary.
                try:
                    text = data[:-3].decode('utf-8')
                except UnicodeDecodeError:
                    continue
            else:
                continue
        if _looks_textual(text):
            return text if not probe else encoding
    return None


def _looks_textual(text: str) -> bool:
    if not text:
        return True
    controls = sum(1 for ch in text[:20000] if ord(ch) < 32 and ch not in '\r\n\t\f')
    return controls / max(len(text[:20000]), 1) < 0.01


def text_encoding(data: bytes) -> str:
    encoding = decode_text(data[:65536], probe=True)
    if encoding is None:
        raise ParseError('unsupported_text_encoding')
    return encoding


# --------------------------------------------------------------------------- PDF


def _pdf(data: bytes, password=None):
    import pypdfium2 as pdfium
    if len(data) > MAX_DOCUMENT_BYTES:
        raise ParseError('document_too_large')
    try:
        return pdfium.PdfDocument(data, password=password)
    except pdfium.PdfiumError as error:
        code = getattr(error, 'err_code', None)
        if code == 4:  # FPDF_ERR_PASSWORD
            raise ParseError('password_required' if password is None else 'password_incorrect') from error
        raise ParseError('corrupt_pdf') from error


def pdf_info(data: bytes, password=None) -> dict:
    pdf = _pdf(data, password)
    try:
        count = len(pdf)
        if count > MAX_PAGES:
            raise ParseError('too_many_pages', str(count))
        pages = []
        for index in range(count):
            page = pdf[index]
            try:
                width, height = page.get_size()
                # pypdfium2 reports the displayed size, already honouring /Rotate.
                pages.append({'page_no': index + 1, 'width_pt': round(width, 3), 'height_pt': round(height, 3), 'rotation': page.get_rotation()})
            finally:
                page.close()
        return {'kind': 'pdf', 'page_count': count, 'pages': pages}
    finally:
        pdf.close()


def _page_to_device(page, width_px, height_px, x, y):
    import ctypes
    import pypdfium2.raw as raw
    dx, dy = ctypes.c_int(), ctypes.c_int()
    raw.FPDF_PageToDevice(page.raw, 0, 0, width_px, height_px, 0, x, y, ctypes.byref(dx), ctypes.byref(dy))
    return dx.value, dy.value


def pdf_pages(data: bytes, first: int, last: int, password=None) -> dict:
    """Native text lines with character-accurate boxes for a page range."""
    import pypdfium2.raw as raw
    pdf = _pdf(data, password)
    try:
        count = len(pdf)
        if not 1 <= first <= last <= count or last - first >= 50:
            raise ParseError('invalid_page_range')
        result = []
        for page_no in range(first, last + 1):
            page = pdf[page_no - 1]
            textpage = page.get_textpage()
            try:
                # Device space of a 10000x10000 virtual raster; normalization keeps
                # boxes independent of any later rendering scale.
                # get_size() is the displayed size (after /Rotate); FPDF_PageToDevice
                # with rotate=0 maps page space into that displayed frame.
                display_w, display_h = page.get_size()
                rotation = page.get_rotation()
                virtual_w = 10000
                virtual_h = max(1, round(10000 * display_h / max(display_w, 1e-6)))
                n = raw.FPDFText_CountChars(textpage.raw)
                lines, current, visible = [], None, 0
                index = 0
                while index < n:
                    code = raw.FPDFText_GetUnicode(textpage.raw, index)
                    if 0xD800 <= code <= 0xDBFF and index + 1 < n:
                        low = raw.FPDFText_GetUnicode(textpage.raw, index + 1)
                        char = chr(0x10000 + ((code - 0xD800) << 10) + (low - 0xDC00))
                        step = 2
                    else:
                        char = chr(code) if code not in (0, 0xFFFE, 0xFFFF) else ''
                        step = 1
                    if char in ('\r', '\n'):
                        if current is not None:
                            lines.append(current)
                            current = None
                        index += step
                        continue
                    if current is None:
                        current = {'text': '', 'box': None}
                    current['text'] += char
                    if char.strip():
                        visible += 1
                        try:
                            left, bottom, right, top = textpage.get_charbox(index)
                        except Exception:
                            left = bottom = right = top = None
                        if left is not None:
                            corners = [_page_to_device(page, virtual_w, virtual_h, x, y) for x, y in ((left, bottom), (right, top))]
                            xs = [c[0] for c in corners]
                            ys = [c[1] for c in corners]
                            box = [min(xs), min(ys), max(xs), max(ys)]
                            if current['box'] is None:
                                current['box'] = box
                            else:
                                b = current['box']
                                current['box'] = [min(b[0], box[0]), min(b[1], box[1]), max(b[2], box[2]), max(b[3], box[3])]
                    index += step
                if current is not None:
                    lines.append(current)
                out_lines = []
                for line in lines:
                    text = line['text'].strip()
                    if not text:
                        continue
                    box = line['box'] or [0, 0, 0, 0]
                    raw_box = [box[0] / virtual_w, box[1] / virtual_h, box[2] / virtual_w, box[3] / virtual_h]
                    clipped = [round(min(1.0, max(0.0, value)), 5) for value in raw_box]
                    item = {'text': text, 'bbox': clipped}
                    if any(abs(a - b) > 0.01 for a, b in zip(raw_box, clipped)):
                        # Text drawn outside the visible page is reported, not hidden.
                        item['outside_page'] = True
                    out_lines.append(item)
                # A footer or a partial text layer says nothing about raster
                # content on the same page. Form XObjects may contain images at
                # arbitrary depth, so treat them conservatively without an
                # unbounded traversal or a depth limit that silently loses text.
                object_types = {obj.type for obj in page.get_objects(filter=[raw.FPDF_PAGEOBJ_IMAGE, raw.FPDF_PAGEOBJ_FORM], max_depth=1)}
                has_raster = raw.FPDF_PAGEOBJ_IMAGE in object_types
                has_forms = raw.FPDF_PAGEOBJ_FORM in object_types
                result.append({'page_no': page_no, 'rotation': rotation, 'display_width_pt': round(display_w, 3), 'display_height_pt': round(display_h, 3),
                               'visible_chars': visible, 'has_raster': has_raster, 'has_forms': has_forms,
                               'needs_ocr': visible < MIN_NATIVE_CHARS or has_raster or has_forms, 'lines': out_lines})
            finally:
                textpage.close()
                page.close()
        return {'pages': result}
    finally:
        pdf.close()


def pdf_render(data: bytes, page_no: int, purpose: str, password=None) -> dict:
    """Render one displayed page to PNG. OCR renders at up to ~200 dpi within budget."""
    pdf = _pdf(data, password)
    try:
        if not 1 <= page_no <= len(pdf):
            raise ParseError('invalid_page_range')
        page = pdf[page_no - 1]
        try:
            width, height = page.get_size()      # displayed size, /Rotate applied
            rotation = page.get_rotation()
            target = 200 / 72 if purpose == 'ocr' else 150 / 72
            budget = OCR_PIXEL_BUDGET if purpose == 'ocr' else 3_000_000
            scale = min(target, (budget / max(width * height, 1)) ** 0.5)
            bitmap = page.render(scale=scale, may_draw_forms=True)
            image = bitmap.to_pil().convert('RGB')
            return _png_result(image, {'scale': round(scale, 6), 'rotation': rotation})
        finally:
            page.close()
    finally:
        pdf.close()


def _png_result(image, extra=None) -> dict:
    import base64
    buffer = io.BytesIO()
    image.save(buffer, format='PNG', optimize=False)
    result = {'png_base64': base64.b64encode(buffer.getvalue()).decode('ascii'), 'width': image.width, 'height': image.height}
    if extra:
        result.update(extra)
    return result


# -------------------------------------------------------------------------- DOCX


def docx_blocks(data: bytes) -> dict:
    """Paragraphs, headings and table cells in document order."""
    check_zip(data)
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    try:
        document = Document(io.BytesIO(data))
    except Exception as error:
        raise ParseError('corrupt_docx') from error
    blocks = []
    table_index = 0
    paragraph_index = 0
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            paragraph_index += 1
            if not text:
                continue
            style = (item.style.name if item.style is not None else '') or ''
            level = None
            match = re.match(r'(?i)^(heading|标题)\s*(\d+)', style)
            if match:
                level = int(match.group(2))
            elif style.lower() == 'title':
                level = 0
            blocks.append({'kind': 'heading' if level is not None else 'paragraph', 'text': text, 'level': level,
                           'locator': {'paragraph': paragraph_index}})
        elif isinstance(item, Table):
            table_index += 1
            # Merged cells repeat the same <w:tc> element. Keep the elements
            # themselves: ids of short-lived lxml proxies can be reused.
            seen = set()
            for row_number, row in enumerate(item.rows, start=1):
                cells = []
                for col_number, cell in enumerate(row.cells, start=1):
                    element = cell._tc
                    if element in seen:
                        continue
                    seen.add(element)
                    text = cell.text.strip()
                    if text:
                        cells.append({'kind': 'table_cell', 'text': text, 'level': None,
                                      'locator': {'table': table_index, 'row': row_number, 'col': col_number}})
                blocks.extend(cells)
        if len(blocks) > 200000:
            raise ParseError('document_too_large')
    return {'kind': 'docx', 'blocks': blocks}


# -------------------------------------------------------------------------- XLSX


def _cell_value(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, float):
        # The shortest round-trip repr preserves typed amounts such as 1234.56.
        return repr(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat(timespec='seconds') if (value.hour or value.minute or value.second) else value.date().isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value)


def xlsx_info(data: bytes) -> dict:
    check_zip(data)
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as error:
        raise ParseError('corrupt_xlsx') from error
    try:
        sheets = []
        total = 0
        for sheet in workbook.worksheets:
            rows, cols = sheet.max_row or 0, sheet.max_column or 0
            total += rows * cols
            sheets.append({'name': sheet.title, 'max_row': rows, 'max_col': cols})
        if total > MAX_SHEET_CELLS:
            raise ParseError('spreadsheet_too_large')
        return {'kind': 'xlsx', 'sheets': sheets}
    finally:
        workbook.close()


def xlsx_rows(data: bytes, sheet_name: str, offset: int, limit: int) -> dict:
    """Cached values with formula markers. Formulas are never evaluated."""
    check_zip(data)
    from openpyxl import load_workbook
    limit = min(limit, MAX_ROWS_PER_CALL)
    values_book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    formula_book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        if sheet_name not in values_book.sheetnames:
            raise ParseError('sheet_not_found')
        values = values_book[sheet_name].iter_rows(min_row=offset + 1, max_row=offset + limit)
        formulas = formula_book[sheet_name].iter_rows(min_row=offset + 1, max_row=offset + limit)
        rows = []
        for row_number, (value_row, formula_row) in enumerate(zip(values, formulas), start=offset + 1):
            cells = []
            for col_number, (value_cell, formula_cell) in enumerate(zip(value_row, formula_row), start=1):
                raw = formula_cell.value
                is_formula = isinstance(raw, str) and raw.startswith('=')
                value = _cell_value(value_cell.value)
                if value is None and not is_formula:
                    cells.append(None)
                    continue
                cells.append({'col': col_number, 'value': value, 'formula': is_formula,
                              'unconfirmed': is_formula and value is None})
            rows.append({'row': row_number, 'cells': cells})
        return {'sheet': sheet_name, 'offset': offset, 'rows': rows}
    finally:
        values_book.close()
        formula_book.close()


# ----------------------------------------------------------------- CSV and text


def delimited_rows(data: bytes, offset: int, limit: int) -> dict:
    if len(data) > MAX_TEXT_BYTES:
        raise ParseError('document_too_large')
    encoding = text_encoding(data)
    text = data.decode(encoding)
    sample = text[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',\t;|')
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ','
    reader = csv.reader(io.StringIO(text, newline=''), delimiter=delimiter)
    limit = min(limit, MAX_ROWS_PER_CALL)
    rows = []
    total = 0
    for row_number, row in enumerate(reader, start=1):
        total = row_number
        if row_number <= offset:
            continue
        if len(rows) < limit:
            rows.append({'row': row_number, 'cells': [cell.strip() for cell in row]})
    return {'kind': 'csv', 'encoding': encoding, 'delimiter': delimiter, 'total_rows': total, 'offset': offset, 'rows': rows}


def text_blocks(data: bytes) -> dict:
    if len(data) > MAX_TEXT_BYTES:
        raise ParseError('document_too_large')
    encoding = text_encoding(data)
    text = data.decode(encoding).replace('\r\n', '\n').replace('\r', '\n')
    blocks = []
    for number, line in enumerate(text.split('\n'), start=1):
        stripped = line.strip()
        if stripped:
            blocks.append({'kind': 'paragraph', 'text': stripped, 'level': None, 'locator': {'line': number}})
    return {'kind': 'text', 'encoding': encoding, 'blocks': blocks}


# ------------------------------------------------------------------------ Images


def image_pages(data: bytes) -> dict:
    """Normalize an image (and TIFF frames) to PNG pages within the OCR budget."""
    from PIL import Image, ImageOps
    Image.MAX_IMAGE_PIXELS = 80_000_000
    if len(data) > 50 * 1024 * 1024:
        raise ParseError('image_too_large')
    try:
        image = Image.open(io.BytesIO(data))
    except Exception as error:
        raise ParseError('corrupt_image') from error
    if image.format not in {'PNG', 'JPEG', 'BMP', 'TIFF', 'WEBP', 'GIF', 'MPO'}:
        raise ParseError('unsupported_image_format', image.format or '')
    pages = []
    frames = getattr(image, 'n_frames', 1) if image.format == 'TIFF' else 1
    for frame in range(min(frames, 200)):
        if frames > 1:
            image.seek(frame)
        exif_orientation = None
        try:
            exif_orientation = image.getexif().get(0x0112)
        except Exception:
            exif_orientation = None
        oriented = ImageOps.exif_transpose(image.copy()) or image.copy()
        rgb = oriented.convert('RGB')
        source_size = [rgb.width, rgb.height]
        scale = min(1.0, (OCR_PIXEL_BUDGET / max(rgb.width * rgb.height, 1)) ** 0.5)
        if scale < 1.0:
            rgb = rgb.resize((max(1, int(rgb.width * scale)), max(1, int(rgb.height * scale))), Image.Resampling.LANCZOS)
        pages.append(_png_result(rgb, {'page_no': frame + 1, 'source_size': source_size, 'scale': round(scale, 6),
                                       'exif_orientation': exif_orientation}))
    return {'kind': 'image', 'format': image.format, 'pages': pages}
